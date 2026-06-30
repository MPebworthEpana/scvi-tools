"""Build MultiVI Zarr changes PowerPoint deck for scvi-tools-2.

Usage (from repo root):
    pip install python-pptx
    python scripts/build_multivi_zarr_deck.py
    python scripts/build_multivi_zarr_deck.py --pdf   # also export PDF (Windows + PowerPoint)

Output:
    docs/presentations/multivi_zarr_changes.pptx
    docs/presentations/multivi_zarr_changes.pdf      (with --pdf)
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PPTX = REPO_ROOT / "docs" / "presentations" / "multivi_zarr_changes.pptx"
OUTPUT_PDF = REPO_ROOT / "docs" / "presentations" / "multivi_zarr_changes.pdf"

# Slide canvas (16:9)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

MARGIN_L = Inches(0.7)
MARGIN_R = Inches(0.7)
CONTENT_W = Inches(11.933)

TITLE_TOP = Inches(0.35)
TITLE_H = Inches(0.9)
BODY_TOP = Inches(1.35)
BODY_H = Inches(5.85)

FONT_TITLE = Pt(30)
FONT_SUBTITLE = Pt(22)
FONT_BODY = Pt(17)
FONT_TABLE = Pt(13)
FONT_NOTE = Pt(12)
FONT_FOOTER = Pt(11)

ROW_H = Inches(0.62)
TABLE_TOP = Inches(1.45)


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _style_paragraph(paragraph, *, size: Pt, bold: bool = False) -> None:
    paragraph.font.size = size
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    paragraph.space_after = Pt(6)


def _add_title_block(slide, title: str) -> None:
    box = slide.shapes.add_textbox(MARGIN_L, TITLE_TOP, CONTENT_W, TITLE_H)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    _style_paragraph(p, size=FONT_TITLE, bold=True)


def _add_bullets(slide, bullets: list[str], *, top=BODY_TOP, height=BODY_H, size=FONT_BODY) -> None:
    box = slide.shapes.add_textbox(MARGIN_L, top, CONTENT_W, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.bullet = True
        _style_paragraph(p, size=size)


def _add_title_slide(prs: Presentation, title: str, subtitle: str, footer: str = "") -> None:
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(MARGIN_L, Inches(2.0), CONTENT_W, Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    _style_paragraph(p, size=Pt(36), bold=True)

    sub = slide.shapes.add_textbox(MARGIN_L, Inches(3.5), CONTENT_W, Inches(0.8))
    stf = sub.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.text = subtitle
    _style_paragraph(sp, size=FONT_SUBTITLE)

    if footer:
        fbox = slide.shapes.add_textbox(MARGIN_L, Inches(6.5), CONTENT_W, Inches(0.4))
        fp = fbox.text_frame.paragraphs[0]
        fp.text = footer
        _style_paragraph(fp, size=FONT_FOOTER)


def _add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = _blank_slide(prs)
    _add_title_block(slide, title)
    _add_bullets(slide, bullets)


def _format_table_cell(cell, text: str, *, bold: bool = False) -> None:
    cell.text = text
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for p in tf.paragraphs:
        _style_paragraph(p, size=FONT_TABLE, bold=bold)


def _add_table_only_slide(
    prs: Presentation,
    title: str,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    slide = _blank_slide(prs)
    _add_title_block(slide, title)

    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_h = ROW_H * n_rows
    table = slide.shapes.add_table(
        n_rows, n_cols, MARGIN_L, TABLE_TOP, CONTENT_W, table_h
    ).table

    col_widths = [CONTENT_W / n_cols] * n_cols
    for j, width in enumerate(col_widths):
        table.columns[j].width = int(width)

    for row in table.rows:
        row.height = int(ROW_H)

    for j, header in enumerate(headers):
        _format_table_cell(table.cell(0, j), header, bold=True)

    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            _format_table_cell(table.cell(i, j), value)


def _add_notes_slide(prs: Presentation, title: str, notes: list[str]) -> None:
    slide = _blank_slide(prs)
    _add_title_block(slide, title)
    _add_bullets(slide, notes, size=FONT_NOTE)


def _add_chart_slide(
    prs: Presentation,
    title: str,
    categories: list[str],
    series_name: str,
    values: list[float],
    notes: list[str],
) -> None:
    slide = _blank_slide(prs)
    _add_title_block(slide, title)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(series_name, values)

    chart_h = Inches(3.6)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        MARGIN_L,
        Inches(1.55),
        CONTENT_W,
        chart_h,
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True

    notes_top = Inches(1.55) + chart_h + Inches(0.25)
    notes_h = SLIDE_H - notes_top - Inches(0.35)
    _add_bullets(slide, notes, top=notes_top, height=notes_h, size=FONT_NOTE)


def build_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _add_title_slide(
        prs,
        "MultiVI at Atlas Scale:\nZarr Streaming in scvi-tools-2",
        "Bounded-memory training + dense ADT lazy loading",
        footer="scvi-tools-2 vs upstream scvi-tools",
    )

    _add_bullet_slide(
        prs,
        "Outline",
        [
            "Two major changes in scvi-tools-2",
            "1. Zarr-backed streaming dataloader",
            "2. Dense ADT zarr.Array with mixed RNA CSR + ADT dense layouts",
            "Agenda: Problem → Architecture → ADT memory → Benchmarks → Usage",
        ],
    )

    _add_bullet_slide(
        prs,
        "The Problem",
        [
            "Atlas-scale MuData (10–20M cells) cannot fit all modality matrices in RAM",
            "RNA: sparse counts work well as zarr-backed CSR",
            "ADT / protein: effectively dense; CSR is a poor fit at this scale",
            "MultiVI training still requires dense minibatches on GPU",
            "Default DataSplitter does not scale the same way for multi-modality zarr MuData",
        ],
    )

    _add_bullet_slide(
        prs,
        "Change 1: Zarr Streaming Architecture",
        [
            "New in scvi-tools-2 (not upstream scvi-tools):",
            "ZarrDataset — IterableDataset with sorted block reads + shuffle buffer",
            "ZarrMultiVIDataModule — Lightning datamodule for MultiVI",
            "MULTIVI.train() auto-selects zarr datamodule for fully zarr-backed MuData",
            "Pipeline: zarr store → block reads → shuffle buffer → minibatch → train",
        ],
    )

    _add_bullet_slide(
        prs,
        "How Streaming Achieves Efficiency",
        [
            "Sorted within-block reads avoid random zarr seeks",
            "Shuffle buffer holds block_size × shuffle_buffer_blocks rows, not full data",
            "Workers reopen zarr via picklable ZarrMatrixSource descriptors",
            "DDP: disjoint per-rank slices with equal batch counts",
            "Defaults: block_size=4096, shuffle_buffer_blocks=16",
        ],
    )

    _add_table_only_slide(
        prs,
        "Change 2: Dense ADT Memory (Table)",
        headers=["Scale", "Full dense ADT in RAM", "Zarr block (4096 rows)"],
        rows=[
            ["10M cells × 300 ADTs", "~12 GB (float32)", "~5 MB per block"],
            ["20M cells × 300 ADTs", "~24 GB (float32)", "~5 MB per block"],
        ],
    )
    _add_notes_slide(
        prs,
        "Change 2: Dense ADT Memory (Notes)",
        [
            "At 10–20M cells, loading full ADT into RAM blows memory",
            "Backed dense zarr.Array lazy-loads only rows for the current block/batch",
            "ADT is dense-ish: CSR adds index overhead without benefit",
            "Use matrix_layout='auto' for RNA CSR + ADT dense on .X",
            "Caveat: from_backed_mudata() defaults to 'csr'; MULTIVI.train auto-path uses 'auto'",
        ],
    )

    _add_bullet_slide(
        prs,
        "What Changed in Code",
        [
            "New: _zarr_dataset.py, _zarr_datamodule.py",
            "MULTIVI.train(): auto zarr datamodule + reload_dataloaders_every_n_epochs=1",
            "zarr.Array support in count validation helpers",
            "Docs: docs/user_guide/use_case/zarr_multivi_streaming.md",
            "Commits: fcd3b31d8, 1c2c96273, 517c16aad",
        ],
    )

    _add_table_only_slide(
        prs,
        "Benchmark: CSR vs Zarr (Results)",
        headers=["Metric", "CSR DataSplitter", "Zarr streaming", "Speedup"],
        rows=[
            ["Loader batches/s (steady)", "13.2", "104.2", "7.9×"],
            ["Training batches/s", "4.6", "7.2", "1.55×"],
            ["Wall time (100 batches)", "21.6 s", "13.9 s", "1.55×"],
            ["First-batch latency", "0.79 s", "1.62 s", "CSR faster"],
        ],
    )
    _add_notes_slide(
        prs,
        "Benchmark: CSR vs Zarr (Setup)",
        [
            "Source: scripts/benchmark_multivi_csr_vs_zarr_wsl_results.json",
            "19,200 cells; 2,000 genes + 2,000 ATAC regions",
            "GPU: RTX 2000 Ada; pin_memory=True; prefetch_to_gpu=True",
            "Zarr wins sustained throughput and total training time",
            "First batch slower due to zarr reopen / worker startup",
        ],
    )

    _add_chart_slide(
        prs,
        "Benchmark Chart: Loader Throughput",
        categories=["CSR DataSplitter", "Zarr streaming"],
        series_name="Batches per second",
        values=[13.2, 104.2],
        notes=[
            "~7.9× higher steady-state loader throughput with ZarrMultiVIDataModule",
            "End-to-end training speedup ~1.55× because GPU compute dominates",
        ],
    )

    _add_table_only_slide(
        prs,
        "Benchmark: Loader Prefetch / Tuning",
        headers=["Config", "Wall (150 batches)", "Loader wait", "Batches/s"],
        rows=[
            ["Baseline (num_workers=0)", "21.9 s", "1.30 s", "6.85"],
            ["Workers + pin_memory", "18.8 s", "0.16 s", "7.96"],
            ["Large batch (512)", "5.9 s", "0.19 s", "25.3"],
        ],
    )
    _add_notes_slide(
        prs,
        "Benchmark: Prefetch Notes",
        [
            "Source: scripts/benchmark_multivi_idle_gap_wsl_results.json",
            "Same ~19k-cell synthetic trimodal MuData on GPU",
            "num_workers + pin_memory sharply reduces GPU idle waiting",
            "Batch size has a large effect on throughput",
        ],
    )

    _add_bullet_slide(
        prs,
        "Positioning vs TileDB-SOMA / CZI",
        [
            "Shared goal: bounded memory + high-throughput streaming",
            "Zarr: MuData zarr stores, MultiVI, mixed CSR + dense modalities",
            "TileDB-SOMA: Census-scale SCVI/SCANVI via SOMA queries",
            "Choose Zarr when atlases are already in mudata.write_zarr format",
        ],
    )

    _add_bullet_slide(
        prs,
        "How to Use",
        [
            "Write MuData to zarr and reopen with backed modality .X",
            "MULTIVI.setup_mudata(...) with correct modality mapping",
            "Mixed RNA + ADT: from_backed_mudata(..., matrix_layout='auto', store_dir=...)",
            "Or model.train(...) on fully zarr-backed MuData (auto datamodule)",
            "Tune: block_size, shuffle_buffer_blocks, num_workers, pin_memory",
        ],
    )

    _add_bullet_slide(
        prs,
        "Limitations",
        [
            "EXPERIMENTAL — MultiVI only",
            "No categorical / continuous covariates yet",
            "CSR modalities densified per batch before GPU (expected)",
            "Windows often needs num_workers=0",
        ],
    )

    _add_bullet_slide(
        prs,
        "Summary & Next Steps",
        [
            "Pillar 1: Zarr streaming for bounded-memory MultiVI training",
            "Pillar 2: Dense ADT zarr.Array for 10–20M-cell lazy loading",
            "Headline benchmarks: 7.9× loader throughput; 1.55× training speedup",
            "Next: 10M+ cell ADT benchmarks, API docs, tutorial notebook",
        ],
    )

    return prs


def export_pdf(pptx_path: Path, pdf_path: Path) -> None:
    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "PDF export requires pywin32 and Microsoft PowerPoint on Windows. "
            "Install with: pip install pywin32"
        ) from exc

    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    powerpoint.Visible = 1
    try:
        presentation = powerpoint.Presentations.Open(str(pptx_path.resolve()), WithWindow=False)
        presentation.SaveAs(str(pdf_path.resolve()), 32)  # ppSaveAsPDF
        presentation.Close()
    finally:
        powerpoint.Quit()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--pdf",
        action="store_true",
        help="Also export PDF via PowerPoint (Windows only)",
    )
    args = parser.parse_args()

    OUTPUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs = build_deck()
    prs.save(OUTPUT_PPTX)
    print(f"Wrote {OUTPUT_PPTX} ({len(prs.slides)} slides)")

    if args.pdf:
        export_pdf(OUTPUT_PPTX, OUTPUT_PDF)
        print(f"Wrote {OUTPUT_PDF}")


if __name__ == "__main__":
    main()
